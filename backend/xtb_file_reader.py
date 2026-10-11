"""Read cash operations only; position sheets are never imported as trades."""
from io import BytesIO
from pathlib import Path
import pandas as pd


def read_xtb_file(file):
    extension = Path(file.filename or '').suffix.lower()
    if extension == '.csv':
        try:
            return pd.read_csv(BytesIO(file.read()), sep=None, engine='python', encoding='utf-8-sig', dtype=object)
        except (pd.errors.ParserError, UnicodeError, ValueError) as error:
            raise ValueError('Nie udało się odczytać CSV. Sprawdź kodowanie UTF-8 i nagłówki.') from error
    if extension != '.xlsx':
        raise ValueError('Wybierz eksport XTB w formacie XLSX lub CSV.')
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(BytesIO(file.read()), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError('Nie udało się odczytać XLSX. Wybierz oryginalny eksport, bez hasła.') from error
    try:
        names = [name for name in workbook.sheetnames if name.strip().casefold() == 'cash operations']
        if len(names) != 1:
            raise ValueError('Eksport XLSX musi zawierać jeden arkusz „Cash Operations”. Innych arkuszy nie importujemy.')
        sheet = workbook[names[0]]
        # Some broker exports declare A1 as the used range despite containing data.
        sheet.reset_dimensions()
        account, header, data, source_rows = None, None, [], []
        for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            labels = [str(value or '').strip().lower() for value in values]
            if header is None:
                if labels[0] == 'account number' and len(values) > 1:
                    account = str(values[1]).strip() if values[1] is not None else None
                if {'type', 'time', 'amount'}.issubset(labels):
                    header = [str(value or '').strip() for value in values]
                    if len(set(labels)) != len(labels) or '' in labels:
                        raise ValueError('Arkusz Cash Operations zawiera powtórzone lub puste nagłówki.')
                elif number >= 100:
                    raise ValueError('Nie znaleziono nagłówków Type, Time i Amount w Cash Operations.')
                continue
            row = list(values[:len(header)]) + [None] * max(0, len(header) - len(values))
            type_index = next(i for i, label in enumerate(header) if label.lower() == 'type')
            if str(row[type_index] or '').strip().lower() in {'total', 'sum', 'suma'}:
                continue
            data.append(row)
            source_rows.append(number)
        if header is None:
            raise ValueError('Nie znaleziono nagłówków Type, Time i Amount w Cash Operations.')
        frame = pd.DataFrame(data, columns=header)
        frame.attrs.update(source_account=account, source_rows=source_rows, cash_operations_xlsx=True)
        return frame
    finally:
        workbook.close()
